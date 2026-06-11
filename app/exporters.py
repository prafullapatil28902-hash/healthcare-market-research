"""
Report exporters.

Generates downloadable artefacts in memory (returned as ``bytes``):
  * Excel workbook   — multi-sheet data + analysis (OpenPyXL)
  * PDF report       — formatted market-research report (ReportLab)
  * Executive summary — PDF and plain-text variants
  * Competitor report — focused PDF on the competitive landscape

All builders accept the cleaned analytics DataFrame so the UI can wire any of
them to a download button.
"""
from __future__ import annotations

import datetime as _dt
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app import competitor_analysis as ca
from app import insights as ins
from app import market_landscape as ml
from app import trend_analysis as ta
from app.executive_summary import ExecutiveSummary, build_executive_summary
from app.validation import ValidationReport

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_BRAND = colors.HexColor("#1f4e79")


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _style_sheet(worksheet) -> None:
    """Apply header styling and auto-ish column widths to an openpyxl sheet."""
    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column_cells in worksheet.columns:
        length = max(
            (len(str(c.value)) for c in column_cells if c.value is not None),
            default=10,
        )
        worksheet.column_dimensions[
            get_column_letter(column_cells[0].column)
        ].width = min(max(length + 2, 12), 40)
    worksheet.freeze_panes = "A2"


def build_excel_report(df: pd.DataFrame) -> bytes:
    """Multi-sheet Excel workbook summarising the full analysis."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        sheets = {
            "Raw Data": df,
            "Company Profiles": ca.company_profiles(df),
            "Market Share": ca.market_share(df),
            "Therapeutic Areas": ml.therapeutic_distribution(df),
            "Regional": ml.regional_distribution(df),
            "Fastest Growing": ta.fastest_growing_areas(df, n=20),
            "Opportunities": ta.market_opportunities(df),
            "Insights": ins.insights_to_dataframe(ins.generate_insights(df)),
        }
        for name, frame in sheets.items():
            safe = frame if frame is not None and not frame.empty else pd.DataFrame(
                {"Info": ["No data"]}
            )
            safe.to_excel(writer, sheet_name=name[:31], index=False)
            _style_sheet(writer.sheets[name[:31]])

    buffer.seek(0)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# PDF helpers
# --------------------------------------------------------------------------- #
def _styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontSize=22,
            textColor=_BRAND,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Subtitle",
            parent=styles["Normal"],
            fontSize=11,
            textColor=colors.grey,
            alignment=TA_CENTER,
            spaceAfter=18,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SectionHeading",
            parent=styles["Heading2"],
            fontSize=14,
            textColor=_BRAND,
            spaceBefore=14,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            parent=styles["Normal"],
            fontSize=10,
            leading=15,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyBullet",
            parent=styles["Normal"],
            fontSize=10,
            leading=15,
            leftIndent=14,
            bulletIndent=4,
            spaceAfter=4,
        )
    )
    return styles


def _df_to_table(df: pd.DataFrame, max_rows: int = 15) -> Table:
    """Convert a DataFrame to a styled ReportLab Table."""
    display = df.head(max_rows).copy()
    # Round floats for presentation.
    for col in display.select_dtypes(include="number").columns:
        display[col] = display[col].map(lambda v: f"{v:,.1f}" if isinstance(v, float) else v)
    data = [list(display.columns)] + display.astype(str).values.tolist()

    table = Table(data, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), _BRAND),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef3f8")]),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _footer(canvas, doc):
    """Page footer with page number and confidentiality note."""
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.grey)
    canvas.drawString(2 * cm, 1 * cm, "Healthcare Market Research Report — Confidential")
    canvas.drawRightString(19 * cm, 1 * cm, f"Page {doc.page}")
    canvas.restoreState()


def _summary_flowables(summary: ExecutiveSummary, styles) -> list:
    """Shared flowables for rendering an ExecutiveSummary into a PDF."""
    flow = [
        Paragraph("Executive Summary", styles["SectionHeading"]),
        Paragraph("<b>Market Overview</b>", styles["Body"]),
        Paragraph(summary.market_overview, styles["Body"]),
        Paragraph("<b>Competitive Landscape</b>", styles["Body"]),
        Paragraph(summary.competitive_landscape, styles["Body"]),
    ]
    for heading, items in (
        ("Key Findings", summary.key_findings),
        ("Opportunities", summary.opportunities),
        ("Recommendations", summary.recommendations),
    ):
        flow.append(Paragraph(f"<b>{heading}</b>", styles["Body"]))
        for item in items:
            flow.append(Paragraph(f"• {item}", styles["BodyBullet"]))
    return flow


# --------------------------------------------------------------------------- #
# PDF reports
# --------------------------------------------------------------------------- #
def build_pdf_report(
    df: pd.DataFrame, validation: ValidationReport | None = None
) -> bytes:
    """Full market-research PDF: summary, KPIs, competitors, areas, insights."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        title="Healthcare Market Research Report",
    )
    styles = _styles()
    today = _dt.date.today().strftime("%d %B %Y")
    flow: list = []

    # --- Cover -------------------------------------------------------------- #
    flow.append(Spacer(1, 4 * cm))
    flow.append(Paragraph("Healthcare Market Research Report", styles["ReportTitle"]))
    flow.append(
        Paragraph(
            "Competitor Analysis &amp; Market Intelligence", styles["Subtitle"]
        )
    )
    flow.append(Paragraph(f"Generated on {today}", styles["Subtitle"]))
    flow.append(PageBreak())

    # --- Executive summary -------------------------------------------------- #
    summary = build_executive_summary(df)
    flow += _summary_flowables(summary, styles)
    flow.append(PageBreak())

    # --- Market KPIs -------------------------------------------------------- #
    kpis = ml.market_kpis(df)
    flow.append(Paragraph("Market Snapshot", styles["SectionHeading"]))
    kpi_table = pd.DataFrame(
        {
            "Metric": [
                "Total Companies",
                "Total Products",
                "Therapeutic Areas",
                "Regions",
                "Total Revenue (USD M)",
                "Average Growth (%)",
            ],
            "Value": [
                kpis["total_companies"],
                kpis["total_products"],
                kpis["total_therapeutic_areas"],
                kpis["total_regions"],
                f"{kpis['total_revenue_usd_m']:,.0f}",
                f"{kpis['avg_growth_pct']:.1f}",
            ],
        }
    )
    flow.append(_df_to_table(kpi_table, max_rows=10))
    flow.append(Spacer(1, 0.5 * cm))

    # --- Competitive landscape --------------------------------------------- #
    flow.append(Paragraph("Top Competitors", styles["SectionHeading"]))
    flow.append(_df_to_table(ca.top_competitors(df, n=10)))
    flow.append(Spacer(1, 0.5 * cm))

    # --- Therapeutic areas -------------------------------------------------- #
    flow.append(Paragraph("Therapeutic Area Landscape", styles["SectionHeading"]))
    flow.append(_df_to_table(ml.therapeutic_distribution(df)))
    flow.append(PageBreak())

    # --- Opportunities & insights ------------------------------------------ #
    flow.append(Paragraph("Market Opportunities", styles["SectionHeading"]))
    flow.append(_df_to_table(ta.market_opportunities(df), max_rows=10))
    flow.append(Spacer(1, 0.5 * cm))

    flow.append(Paragraph("Analyst Insights", styles["SectionHeading"]))
    for insight in ins.generate_insights(df):
        flow.append(
            Paragraph(f"<b>{insight.category}:</b> {insight.text}", styles["BodyBullet"])
        )

    # --- Validation appendix ------------------------------------------------ #
    if validation is not None:
        flow.append(PageBreak())
        flow.append(Paragraph("Data Validation Report", styles["SectionHeading"]))
        flow.append(
            Paragraph(
                f"Quality score: <b>{validation.quality_score}/100</b> — "
                f"{validation.error_count} error(s), "
                f"{validation.warning_count} warning(s) across "
                f"{validation.total_rows} rows.",
                styles["Body"],
            )
        )
        vdf = validation.to_dataframe()
        if not vdf.empty:
            flow.append(_df_to_table(vdf, max_rows=20))
        else:
            flow.append(Paragraph("No data-quality issues detected.", styles["Body"]))

    doc.build(flow, onFirstPage=_footer, onLaterPages=_footer)
    buffer.seek(0)
    return buffer.getvalue()


def build_executive_summary_pdf(df: pd.DataFrame) -> bytes:
    """Standalone executive-summary PDF."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title="Executive Summary")
    styles = _styles()
    summary = build_executive_summary(df)

    flow = [
        Paragraph(summary.title, styles["ReportTitle"]),
        Paragraph(f"Generated on {summary.generated_on}", styles["Subtitle"]),
    ]
    flow += _summary_flowables(summary, styles)
    doc.build(flow, onFirstPage=_footer, onLaterPages=_footer)
    buffer.seek(0)
    return buffer.getvalue()


def build_competitor_pdf(df: pd.DataFrame) -> bytes:
    """Focused competitor-analysis PDF."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title="Competitor Analysis")
    styles = _styles()
    today = _dt.date.today().strftime("%d %B %Y")

    flow = [
        Paragraph("Competitor Analysis Report", styles["ReportTitle"]),
        Paragraph(f"Generated on {today}", styles["Subtitle"]),
        Paragraph("Company Profiles", styles["SectionHeading"]),
        _df_to_table(ca.company_profiles(df), max_rows=20),
        Spacer(1, 0.5 * cm),
        Paragraph("Market Share", styles["SectionHeading"]),
        _df_to_table(ca.market_share(df), max_rows=20),
        Spacer(1, 0.5 * cm),
        Paragraph("Therapeutic Area Coverage", styles["SectionHeading"]),
        _df_to_table(
            ca.therapeutic_coverage(df).reset_index(), max_rows=20
        ),
    ]
    doc.build(flow, onFirstPage=_footer, onLaterPages=_footer)
    buffer.seek(0)
    return buffer.getvalue()


def build_executive_summary_text(df: pd.DataFrame) -> bytes:
    """Plain-text executive summary (.txt)."""
    summary = build_executive_summary(df)
    return summary.to_plain_text().encode("utf-8")
